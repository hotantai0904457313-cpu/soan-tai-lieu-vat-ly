"""
Trích hình ảnh trong file Word bài tập → làm nét → xuất file Word mới
chỉ gồm "Câu N." + hình của câu đó (bỏ toàn bộ nội dung chữ của câu).

Luồng:
  1. Quét các đoạn văn (KHÔNG quét ô bảng) theo đúng thứ tự tài liệu,
     lấy ảnh + kích thước hiển thị gốc (wp:extent) + văn bản ngữ cảnh.
  2. Gán số câu cho từng ảnh: regex cục bộ trước, sau đó nhờ AI
     (9Router → Gemini) rà lại toàn bộ để bắt các ca ảnh nằm cách xa
     dòng "Câu N." hoặc đề đánh số kiểu khác.
  3. Làm nét ảnh bằng Pillow (phóng to + unsharp + tăng tương phản).
     KHÔNG dùng AI sinh ảnh — hình vật lý phải giữ nguyên nội dung.
  4. Dựng file .docx mới: "Câu N." in đậm + hình bên dưới.
"""
import io
import json
import os
import re
import uuid
from typing import List, Optional

from PIL import Image as PILImage, ImageEnhance, ImageFilter, ImageOps

TMP_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)),
                       'data', 'uploads', 'docx_images')
os.makedirs(TMP_DIR, exist_ok=True)

# ── Namespace trong OOXML ────────────────────────────────────────────
_NS_R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
_TAG_BLIP = '{http://schemas.openxmlformats.org/drawingml/2006/main}blip'
_TAG_EXTENT = ('{http://schemas.openxmlformats.org/drawingml/2006/'
               'wordprocessingDrawing}extent')
_TAG_VML = '{urn:schemas-microsoft-com:vml}imagedata'
_EMU_PER_CM = 360000.0

# Nhãn câu: "Câu 12", "Bài 3.", "12." , "12)" — KHÔNG bắt "1.5 kg" (số thập phân)
_Q_LABEL_RE = re.compile(
    r'^\s*(?:'
    r'(?:c[âa]u|b[àa]i)\s*[.:]?\s*(\d{1,3})'
    r'|(\d{1,3})\s*[.)](?!\d)'
    r')',
    re.IGNORECASE)

# Mức làm nét: (ngưỡng phóng to, hệ số phóng, radius, percent, threshold,
#               cutoff autocontrast, hệ số tương phản)
_LEVELS = {
    'nhe':  (500, 2, 1.5, 120, 3, 0.0, 1.08),
    'vua':  (600, 2, 2.0, 165, 3, 0.5, 1.18),
    'manh': (900, 3, 2.5, 220, 2, 1.0, 1.35),
}
_MAX_SIDE = 2600          # chặn ảnh phình quá to sau khi phóng
_CTX_CHARS = 160          # số ký tự ngữ cảnh gửi cho AI mỗi ảnh


def sharpen_bytes(img_bytes: bytes, level: str = 'manh') -> Optional[PILImage.Image]:
    """Làm nét 1 ảnh. Trả None nếu Pillow không đọc được (EMF/WMF...)."""
    thr, scale, radius, percent, threshold, cutoff, contrast = \
        _LEVELS.get(level, _LEVELS['manh'])
    try:
        img = PILImage.open(io.BytesIO(img_bytes))
        img.load()
    except Exception:
        return None

    # Nền trong suốt → ghép lên nền trắng (Word in ra nền trắng)
    if img.mode in ('RGBA', 'LA', 'P'):
        img = img.convert('RGBA')
        bg = PILImage.new('RGBA', img.size, (255, 255, 255, 255))
        img = PILImage.alpha_composite(bg, img).convert('RGB')
    elif img.mode != 'RGB':
        img = img.convert('RGB')

    # Phóng to ảnh nhỏ trước khi làm nét (nét hơn khi in, không răng cưa)
    w, h = img.size
    if min(w, h) < thr:
        f = min(scale, _MAX_SIDE / max(w, h)) if max(w, h) else 1
        if f > 1:
            img = img.resize((int(w * f), int(h * f)), PILImage.LANCZOS)

    img = img.filter(ImageFilter.UnsharpMask(
        radius=radius, percent=percent, threshold=threshold))
    if cutoff:
        img = ImageOps.autocontrast(img, cutoff=cutoff)
    if contrast != 1.0:
        img = ImageEnhance.Contrast(img).enhance(contrast)
    return img


def _para_text(para) -> str:
    return (para.text or '').strip()


def extract_paragraph_images(docx_path: str) -> tuple:
    """Quét đoạn văn top-level (bỏ ô bảng) → danh sách ảnh theo thứ tự.

    Trả (items, n_in_tables) — mỗi item:
      {'blob', 'w_cm', 'h_cm', 'ctx', 'q_local'}
    """
    from docx import Document as DocxDoc

    doc = DocxDoc(docx_path)
    part = doc.part
    items: List[dict] = []
    seen_rids_body = set()

    cur_q: Optional[int] = None
    recent_text: List[str] = []

    for para in doc.paragraphs:
        text = _para_text(para)
        if text:
            m = _Q_LABEL_RE.match(text)
            if m:
                cur_q = int(m.group(1) or m.group(2))
            recent_text.append(text)
            recent_text = recent_text[-3:]

        last_extent = None
        for elem in para._p.iter():
            if elem.tag == _TAG_EXTENT:
                last_extent = elem
                continue
            rid = None
            if elem.tag == _TAG_BLIP:
                rid = elem.get(_NS_R + 'embed')
            elif elem.tag == _TAG_VML:
                rid = elem.get(_NS_R + 'id')
            if not rid:
                continue
            seen_rids_body.add(rid)
            try:
                blob = part.related_parts[rid].blob
            except Exception:
                continue
            w_cm = h_cm = None
            if last_extent is not None:
                try:
                    w_cm = int(last_extent.get('cx')) / _EMU_PER_CM
                    h_cm = int(last_extent.get('cy')) / _EMU_PER_CM
                except (TypeError, ValueError):
                    w_cm = h_cm = None
            ctx = ' '.join(recent_text)[-_CTX_CHARS:]
            items.append({'blob': blob, 'w_cm': w_cm, 'h_cm': h_cm,
                          'ctx': ctx, 'q_local': cur_q})

    # Đếm ảnh nằm trong bảng (không lấy, nhưng phải báo để thầy biết)
    n_in_tables = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for elem in para._p.iter():
                        if elem.tag == _TAG_BLIP or elem.tag == _TAG_VML:
                            n_in_tables += 1
    return items, n_in_tables


_AI_PROMPT = """Bạn đang xử lý một file bài tập Vật lý. Dưới đây là danh sách các HÌNH ẢNH trong file, kèm đoạn văn bản đứng ngay trước mỗi hình.

Nhiệm vụ: xác định mỗi hình thuộc CÂU SỐ MẤY.

Quy tắc:
- Chỉ trả về số thứ tự câu (số nguyên). Nếu không xác định được, trả null.
- Hình thường thuộc về câu được nhắc gần nhất phía trước nó.
- Nhiều hình liên tiếp có thể cùng thuộc một câu.
- Chỉ dựa vào văn bản được cung cấp, không suy đoán thêm.

Trả về DUY NHẤT một JSON object dạng {"1": 3, "2": 3, "3": 7, "4": null}
với khóa là số thứ tự hình (bắt đầu từ 1), giá trị là số câu.

DANH SÁCH HÌNH:
"""


def _parse_ai_map(raw: str, n: int) -> dict:
    """Đọc JSON AI trả về → {chỉ số ảnh (1-based): số câu}."""
    if not raw:
        return {}
    s = raw.strip()
    s = re.sub(r'^```(?:json)?\s*|\s*```$', '', s, flags=re.MULTILINE).strip()
    data = None
    try:
        data = json.loads(s)
    except Exception:
        m = re.search(r'\{.*\}', s, re.DOTALL)
        if m:
            try:
                data = json.loads(m.group(0))
            except Exception:
                data = None
    if not isinstance(data, dict):
        return {}
    out = {}
    for k, v in data.items():
        try:
            idx = int(str(k).strip())
        except ValueError:
            continue
        if not (1 <= idx <= n):
            continue
        if v is None:
            continue
        try:
            num = int(v)
        except (TypeError, ValueError):
            m = re.search(r'\d+', str(v))
            if not m:
                continue
            num = int(m.group(0))
        if 0 < num <= 999:
            out[idx] = num
    return out


def assign_questions(items: List[dict], gemini_key: str = '',
                     niner_url: str = '', niner_key: str = '') -> str:
    """Gán 'q' cho từng ảnh. Ưu tiên AI, thiếu chỗ nào lấy regex bù.
    Trả về tên phương pháp đã dùng: 'ai' | 'regex'."""
    for it in items:
        it['q'] = it.get('q_local')

    if not items or not (gemini_key or niner_key):
        return 'regex'

    lines = []
    for i, it in enumerate(items, 1):
        ctx = (it.get('ctx') or '').replace('\n', ' ').strip() or '(không có chữ)'
        lines.append(f'Hình {i}: "{ctx}"')
    prompt = _AI_PROMPT + '\n'.join(lines) + '\n\nJSON:'

    try:
        from core.pdf_to_word import _call_ai_fix
        raw = _call_ai_fix(prompt, gemini_key, niner_url, niner_key)
    except Exception:
        raw = ''
    mapping = _parse_ai_map(raw, len(items))
    if not mapping:
        return 'regex'
    for i, it in enumerate(items, 1):
        if i in mapping:
            it['q'] = mapping[i]
    return 'ai'


def build_images_docx(items: List[dict], out_path: str, title: str,
                      level: str = 'manh') -> dict:
    """Làm nét + dựng file Word chỉ gồm 'Câu N.' + hình. Trả thống kê."""
    from docx import Document as DocxDoc
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    # Gom theo số câu, giữ thứ tự xuất hiện trong file gốc
    groups: dict = {}
    for it in items:
        groups.setdefault(it.get('q'), []).append(it)
    known = sorted(k for k in groups if k is not None)
    order = known + ([None] if None in groups else [])

    doc = DocxDoc()
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(13)
    st.paragraph_format.space_after = Pt(2)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        sec.left_margin = sec.right_margin = Cm(1.5)
    usable_cm = 21.0 - 3.0     # A4 rộng 21cm, trừ lề trái+phải

    h = doc.add_paragraph()
    hr = h.add_run(title)
    hr.bold = True
    hr.font.size = Pt(15)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_after = Pt(10)

    n_ok = n_bad = 0
    tmp_files = []
    for q in order:
        label = f'Câu {q}.' if q is not None else 'Hình chưa xác định câu'
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        run = p.add_run(label)
        run.bold = True

        for it in groups[q]:
            img = sharpen_bytes(it['blob'], level)
            if img is None:
                n_bad += 1
                continue
            tmp = os.path.join(TMP_DIR, f'{uuid.uuid4().hex[:10]}.png')
            img.save(tmp, 'PNG')
            tmp_files.append(tmp)

            # Giữ đúng bề rộng hiển thị như file gốc (ảnh nét hơn vì nhiều điểm
            # ảnh hơn trên cùng diện tích); không có extent thì suy từ 96 DPI.
            w_cm = it.get('w_cm')
            if not w_cm or w_cm <= 0:
                w_cm = img.size[0] / 96 * 2.54 / _LEVELS.get(
                    level, _LEVELS['manh'])[1]
            w_cm = max(2.0, min(w_cm, usable_cm))

            ip = doc.add_paragraph()
            ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ip.add_run().add_picture(tmp, width=Cm(w_cm))
            n_ok += 1

    doc.save(out_path)
    for f in tmp_files:
        try:
            os.remove(f)
        except OSError:
            pass

    return {'so_cau': len([k for k in order if k is not None]),
            'so_hinh': n_ok, 'hinh_loi': n_bad}


def extract_question_images(docx_path: str, out_path: str, title: str,
                            level: str = 'manh', gemini_key: str = '',
                            niner_url: str = '', niner_key: str = '') -> dict:
    """Đầu vào .docx bài tập → .docx chỉ gồm số câu + hình đã làm nét."""
    items, n_tables = extract_paragraph_images(docx_path)
    if not items:
        raise ValueError('Không tìm thấy hình nào trong file (ngoài bảng). '
                         'Nếu hình nằm trong bảng, hãy báo để bật chế độ quét bảng.')
    method = assign_questions(items, gemini_key, niner_url, niner_key)
    stats = build_images_docx(items, out_path, title, level)
    stats['hinh_trong_bang'] = n_tables
    stats['phuong_phap'] = method
    return stats
